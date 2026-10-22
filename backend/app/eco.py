"""SMART HEALTHCARE ECOSYSTEM — extended module routers.

Modules: appointments, blood bank, emergency response, AI health assistant,
articles, rural services, nutrition, wellness, medicines, billing, beds,
global search, user settings, ecosystem analytics, risk-ML lab, extra reports.
All data synthetic/demo. Educational + decision-support only.
"""
import csv
import io
import json
from datetime import date, datetime, timedelta
from typing import List, Optional

from fastapi import APIRouter, Depends, HTTPException, Query
from fastapi.responses import Response
from pydantic import BaseModel
from sqlalchemy import func
from sqlalchemy.orm import Session, joinedload

from .database import get_db
from .auth import get_current_user, require_roles
from .live import hub
from .models import (
    User, Patient, Doctor, Department, Visit, Appointment, DoctorSchedule,
    BloodDonor, BloodInventory, BloodRequest, Ambulance, EmergencyRequest,
    HealthArticle, ChatSession, ChatMessage, NutritionResource, WellnessLog,
    Medicine, ServiceDir, Invoice, Bed, MlPrediction, UserSetting, Notification, AuditLog,
)
from .kb import chat_reply, ARTICLES, MEALS

router = APIRouter()
GROUPS = ["A+", "A-", "B+", "B-", "AB+", "AB-", "O+", "O-"]


def _audit(db: Session, user_id, action, entity="", entity_id=None, details=""):
    db.add(AuditLog(user_id=user_id, action=action, entity=entity, entity_id=entity_id, details=str(details)[:2000]))


def _notify(db: Session, type: str, title: str, message: str, role_target="", user_id=None, related_id=None):
    n = Notification(type=type, title=title, message=message, role_target=role_target, user_id=user_id, related_id=related_id)
    db.add(n); db.flush()
    return n


def _export(rows: List[dict], columns: List[str], title: str, fmt: str):
    if fmt == "csv":
        buf = io.StringIO()
        w = csv.DictWriter(buf, fieldnames=columns)
        w.writeheader()
        for r in rows:
            w.writerow({c: r.get(c, "") for c in columns})
        return Response(content=buf.getvalue(), media_type="text/csv",
                        headers={"Content-Disposition": f"attachment; filename={title}.csv"})
    if fmt == "excel":
        from openpyxl import Workbook
        wb = Workbook(); ws = wb.active; ws.title = title[:30]
        ws.append(columns)
        for r in rows:
            ws.append([r.get(c, "") for c in columns])
        buf = io.BytesIO(); wb.save(buf)
        return Response(content=buf.getvalue(), media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                        headers={"Content-Disposition": f"attachment; filename={title}.xlsx"})
    if fmt == "pdf":
        from fpdf import FPDF
        def _safe(s):
            return str(s).replace("\u2014", "-").encode("latin-1", "replace").decode("latin-1")
        pdf = FPDF(orientation="L"); pdf.add_page(); pdf.set_font("Helvetica", "B", 13)
        pdf.cell(0, 10, _safe(title + " (Synthetic Demo Data)"), new_x="LMARGIN", new_y="NEXT")
        pdf.set_font("Helvetica", "", 8)
        for r in rows[:100]:
            pdf.cell(0, 5, _safe(" | ".join(f"{c}: {str(r.get(c,''))[:24]}" for c in columns[:6])), new_x="LMARGIN", new_y="NEXT")
        return Response(content=bytes(pdf.output()), media_type="application/pdf",
                        headers={"Content-Disposition": f"attachment; filename={title}.pdf"})
    return {"title": title, "count": len(rows), "items": rows, "synthetic_data": True}


# ================= APPOINTMENTS =================
class ApptIn(BaseModel):
    patient_id: int
    doctor_id: Optional[int] = None
    department_id: Optional[int] = None
    date: date
    time_slot: str = ""
    reason: str = ""


def _appt(a: Appointment) -> dict:
    return {"id": a.id, "patient_id": a.patient_id, "patient_name": a.patient.full_name if a.patient else "",
            "patient_code": a.patient.patient_code if a.patient else "", "doctor_id": a.doctor_id,
            "doctor_name": a.doctor.full_name if a.doctor else "", "department_id": a.department_id,
            "department": a.department.name if a.department else "", "date": a.date.isoformat() if a.date else "",
            "time_slot": a.time_slot, "reason": a.reason, "status": a.status, "notes": a.notes,
            "created_at": a.created_at.isoformat() if a.created_at else None}


@router.get("/api/appointments")
def list_appts(status: str = "", doctor_id: Optional[int] = None, date_: str = Query("", alias="date"),
               upcoming: bool = False, db: Session = Depends(get_db), user: User = Depends(get_current_user)):
    q = db.query(Appointment).options(joinedload(Appointment.patient), joinedload(Appointment.doctor), joinedload(Appointment.department))
    if user.role == "patient":
        myp = [p.id for p in db.query(Patient).filter_by(user_id=user.id).all()]
        q = q.filter(Appointment.patient_id.in_(myp or [-1]))
    if user.role == "doctor":
        mydoc = db.query(Doctor).filter(Doctor.full_name.ilike(f"%{user.full_name.replace('Dr.','').strip().split(' ')[-1]}%")).first()
        doc_ids = [d.id for d in db.query(Doctor).all() if user.username in (d.license_no or "")] or ([mydoc.id] if mydoc else [])
        # fallback: show all to doctor role demo (assigned filter when linked)
        if doc_ids:
            q = q.filter(Appointment.doctor_id.in_(doc_ids))
    if status:
        q = q.filter_by(status=status)
    if doctor_id:
        q = q.filter_by(doctor_id=doctor_id)
    if date_:
        q = q.filter(Appointment.date == date.fromisoformat(date_))
    if upcoming:
        q = q.filter(Appointment.date >= date.today(), Appointment.status.in_(["requested", "confirmed", "rescheduled"]))
    rows = q.order_by(Appointment.date.desc(), Appointment.id.desc()).limit(500).all()
    return [_appt(a) for a in rows]


@router.post("/api/appointments")
async def create_appt(body: ApptIn, db: Session = Depends(get_db), user: User = Depends(get_current_user)):
    p = db.get(Patient, body.patient_id)
    if not p:
        raise HTTPException(404, "Patient not found")
    if user.role == "patient" and p.user_id != user.id:
        raise HTTPException(403, "You can only book appointments for your own profile")
    # prevent double-booking same doctor/date/slot
    if body.doctor_id and body.time_slot:
        clash = db.query(Appointment).filter_by(doctor_id=body.doctor_id, date=body.date, time_slot=body.time_slot).filter(Appointment.status.in_(["requested", "confirmed", "rescheduled"])).first()
        if clash:
            raise HTTPException(400, "That slot is already booked for this doctor")
    a = Appointment(**body.model_dump(), status="requested", created_by=user.id)
    db.add(a); db.flush()
    _audit(db, user.id, "create", "appointment", a.id, f"patient {body.patient_id} {body.date} {body.time_slot}")
    _notify(db, "appointment", "New appointment request", f"#{a.id} — {body.date} {body.time_slot} {body.reason[:60]}", role_target="doctor", related_id=a.id)
    db.commit(); db.refresh(a)
    await hub.broadcast("appointment_updated", {"reason": "created", "appointment": _appt(a)})
    return _appt(a)


class ApptPatch(BaseModel):
    status: Optional[str] = None
    date: Optional[date] = None
    time_slot: Optional[str] = None
    notes: Optional[str] = None


@router.patch("/api/appointments/{aid}")
async def patch_appt(aid: int, body: ApptPatch, db: Session = Depends(get_db), user: User = Depends(require_roles("admin", "doctor", "triage"))):
    a = db.get(Appointment, aid)
    if not a:
        raise HTTPException(404, "Appointment not found")
    if body.status:
        if body.status not in ("requested", "confirmed", "rejected", "completed", "cancelled", "rescheduled"):
            raise HTTPException(400, "Invalid status")
        a.status = body.status
    if body.date:
        a.date = body.date
    if body.time_slot is not None:
        a.time_slot = body.time_slot
    if body.notes is not None:
        a.notes = body.notes
    _audit(db, user.id, "update", "appointment", aid, body.model_dump_json()[:400])
    p = db.get(Patient, a.patient_id)
    _notify(db, "appointment", f"Appointment {a.status}", f"#{aid} on {a.date} {a.time_slot} — {a.status}", role_target="", user_id=p.user_id if p else None, related_id=aid)
    db.commit()
    await hub.broadcast("appointment_updated", {"reason": "updated", "appointment": _appt(a)})
    return _appt(a)


@router.post("/api/appointments/{aid}/cancel")
async def cancel_appt(aid: int, db: Session = Depends(get_db), user: User = Depends(get_current_user)):
    a = db.get(Appointment, aid)
    if not a:
        raise HTTPException(404, "Appointment not found")
    if user.role == "patient":
        p = db.get(Patient, a.patient_id)
        if not p or p.user_id != user.id:
            raise HTTPException(403, "Not your appointment")
    a.status = "cancelled"
    _audit(db, user.id, "cancel", "appointment", aid, "")
    db.commit()
    await hub.broadcast("appointment_updated", {"reason": "cancelled", "appointment": _appt(a)})
    return {"ok": True}


@router.get("/api/appointments/stats")
def appt_stats(db: Session = Depends(get_db), user: User = Depends(get_current_user)):
    today = date.today()
    by_status = {s: db.query(Appointment).filter_by(status=s).count() for s in ("requested", "confirmed", "rejected", "completed", "cancelled")}
    today_n = db.query(Appointment).filter_by(date=today).count()
    daily = []
    for i in range(6, -1, -1):
        d = today - timedelta(days=i)
        daily.append({"date": d.isoformat(), "count": db.query(Appointment).filter_by(date=d).count()})
    upcoming = db.query(Appointment).filter(Appointment.date >= today, Appointment.status.in_(["requested", "confirmed", "rescheduled"])).count()
    return {"today": today_n, "upcoming": upcoming, "by_status": by_status, "daily": daily}


class SchedIn(BaseModel):
    doctor_id: int
    weekday: int
    start_time: str = "09:00"
    end_time: str = "13:00"
    slot_minutes: int = 20
    active: bool = True


@router.get("/api/schedules")
def list_sched(doctor_id: Optional[int] = None, db: Session = Depends(get_db), user: User = Depends(get_current_user)):
    q = db.query(DoctorSchedule)
    if doctor_id:
        q = q.filter_by(doctor_id=doctor_id)
    return [{"id": s.id, "doctor_id": s.doctor_id, "weekday": s.weekday, "start_time": s.start_time, "end_time": s.end_time, "slot_minutes": s.slot_minutes, "active": s.active} for s in q.order_by(DoctorSchedule.weekday).limit(500)]


@router.post("/api/schedules")
def create_sched(body: SchedIn, db: Session = Depends(get_db), user: User = Depends(require_roles("admin", "doctor"))):
    s = DoctorSchedule(**body.model_dump())
    db.add(s); db.commit(); db.refresh(s)
    return {"id": s.id}


@router.delete("/api/schedules/{sid}")
def del_sched(sid: int, db: Session = Depends(get_db), user: User = Depends(require_roles("admin", "doctor"))):
    s = db.get(DoctorSchedule, sid)
    if s:
        db.delete(s); db.commit()
    return {"ok": True}


@router.get("/api/doctors/{did}/slots")
def doctor_slots(did: int, date_: str = Query(..., alias="date"), db: Session = Depends(get_db), user: User = Depends(get_current_user)):
    d = date.fromisoformat(date_)
    scheds = db.query(DoctorSchedule).filter_by(doctor_id=did, weekday=d.weekday(), active=True).all()
    if not scheds:
        return {"date": date_, "slots": [], "note": "Doctor has no schedule this weekday"}
    booked = {a.time_slot for a in db.query(Appointment).filter_by(doctor_id=did, date=d).filter(Appointment.status.in_(["requested", "confirmed", "rescheduled"])).all()}
    slots = []
    for s in scheds:
        try:
            sh, sm = map(int, s.start_time.split(":"))
            eh, em = map(int, s.end_time.split(":"))
        except Exception:
            continue
        cur = sh * 60 + sm
        end = eh * 60 + em
        while cur + s.slot_minutes <= end:
            label = f"{cur//60:02d}:{cur%60:02d}"
            slots.append({"time": label, "booked": label in booked})
            cur += s.slot_minutes
    return {"date": date_, "slots": slots}


# ================= BLOOD BANK =================
@router.get("/api/blood/donors")
def list_donors(search: str = "", group: str = "", db: Session = Depends(get_db), user: User = Depends(get_current_user)):
    q = db.query(BloodDonor)
    if search:
        like = f"%{search}%"
        q = q.filter((BloodDonor.full_name.ilike(like)) | (BloodDonor.phone.ilike(like)))
    if group:
        q = q.filter_by(blood_group=group)
    rows = q.order_by(BloodDonor.id.desc()).limit(300).all()
    return [{"id": x.id, "full_name": x.full_name, "blood_group": x.blood_group, "phone": x.phone, "gender": x.gender, "age": x.age, "last_donation": x.last_donation.isoformat() if x.last_donation else None, "donations_count": x.donations_count} for x in rows]


class DonorIn(BaseModel):
    full_name: str
    blood_group: str
    phone: str = ""
    gender: str = ""
    age: int = 0


@router.post("/api/blood/donors")
def create_donor(body: DonorIn, db: Session = Depends(get_db), user: User = Depends(require_roles("admin", "blood", "triage"))):
    if body.blood_group not in GROUPS:
        raise HTTPException(400, "Invalid blood group")
    x = BloodDonor(**body.model_dump())
    db.add(x)
    _audit(db, user.id, "create", "donor", None, body.full_name)
    db.commit(); db.refresh(x)
    return {"id": x.id}


@router.get("/api/blood/inventory")
def blood_inventory(db: Session = Depends(get_db), user: User = Depends(get_current_user)):
    rows = {r.blood_group: r for r in db.query(BloodInventory).all()}
    return [{"group": g, "available": rows[g].units_available if g in rows else 0, "reserved": rows[g].units_reserved if g in rows else 0} for g in GROUPS]


class InvPatch(BaseModel):
    units_available: Optional[int] = None
    units_reserved: Optional[int] = None


@router.patch("/api/blood/inventory/{group}")
async def patch_inventory(group: str, body: InvPatch, db: Session = Depends(get_db), user: User = Depends(require_roles("admin", "blood"))):
    r = db.query(BloodInventory).filter_by(blood_group=group).first()
    if not r:
        raise HTTPException(404, "Group not found")
    if body.units_available is not None:
        r.units_available = max(0, body.units_available)
    if body.units_reserved is not None:
        r.units_reserved = max(0, body.units_reserved)
    _audit(db, user.id, "update", "blood_inventory", r.id, f"{group} avail={r.units_available}")
    if r.units_available < 5:
        _notify(db, "resource", f"Low blood stock: {group}", f"Only {r.units_available} units of {group} available", role_target="blood", related_id=r.id)
    db.commit()
    await hub.broadcast("blood_updated", {"group": group})
    return {"ok": True}


@router.get("/api/blood/requests")
def list_blood_reqs(status: str = "", db: Session = Depends(get_db), user: User = Depends(get_current_user)):
    q = db.query(BloodRequest)
    if status:
        q = q.filter_by(status=status)
    rows = q.order_by(BloodRequest.id.desc()).limit(300).all()
    return [{"id": r.id, "hospital_unit": r.hospital_unit, "blood_group": r.blood_group, "units": r.units, "priority": r.priority, "status": r.status, "created_at": r.created_at.isoformat() if r.created_at else None} for r in rows]


class BReqIn(BaseModel):
    hospital_unit: str = ""
    blood_group: str
    units: int = 1
    priority: str = "routine"


@router.post("/api/blood/requests")
async def create_blood_req(body: BReqIn, db: Session = Depends(get_db), user: User = Depends(get_current_user)):
    if body.blood_group not in GROUPS:
        raise HTTPException(400, "Invalid blood group")
    r = BloodRequest(**body.model_dump(), requested_by=user.id)
    db.add(r); db.flush()
    _notify(db, "info", "New blood request", f"{body.units} units of {body.blood_group} for {body.hospital_unit}", role_target="blood", related_id=r.id)
    db.commit(); db.refresh(r)
    await hub.broadcast("blood_updated", {"reason": "request"})
    return {"id": r.id}


class BReqPatch(BaseModel):
    status: str


@router.patch("/api/blood/requests/{rid}")
async def patch_blood_req(rid: int, body: BReqPatch, db: Session = Depends(get_db), user: User = Depends(require_roles("admin", "blood"))):
    r = db.get(BloodRequest, rid)
    if not r:
        raise HTTPException(404, "Request not found")
    if body.status not in ("requested", "approved", "issued", "rejected", "fulfilled", "cancelled"):
        raise HTTPException(400, "Invalid status")
    inv = db.query(BloodInventory).filter_by(blood_group=r.blood_group).first()
    if body.status == "issued" and inv:
        if inv.units_available < r.units:
            raise HTTPException(400, f"Insufficient stock: {inv.units_available} units of {r.blood_group}")
        inv.units_available -= r.units
    r.status = body.status
    _audit(db, user.id, "update", "blood_request", rid, body.status)
    db.commit()
    await hub.broadcast("blood_updated", {"reason": "status"})
    return {"ok": True}


# ================= EMERGENCY =================
def _er(r: EmergencyRequest) -> dict:
    try:
        tl = json.loads(r.timeline or "[]")
    except Exception:
        tl = []
    return {"id": r.id, "patient_name": r.patient_name, "phone": r.phone, "emergency_type": r.emergency_type,
            "location": r.location, "priority": r.priority, "status": r.status, "ambulance_id": r.ambulance_id,
            "ambulance": r.ambulance.vehicle_no if r.ambulance else "", "assigned_staff": r.assigned_staff,
            "notes": r.notes, "timeline": tl, "created_at": r.created_at.isoformat() if r.created_at else None}


class ERIn(BaseModel):
    patient_name: str = ""
    phone: str = ""
    emergency_type: str = "medical"
    location: str = ""
    priority: str = "urgent"
    notes: str = ""


@router.get("/api/emergency/requests")
def list_er(status: str = "", db: Session = Depends(get_db), user: User = Depends(get_current_user)):
    q = db.query(EmergencyRequest).options(joinedload(EmergencyRequest.ambulance))
    if status:
        q = q.filter_by(status=status)
    return [_er(r) for r in q.order_by(EmergencyRequest.id.desc()).limit(300).all()]


@router.post("/api/emergency/requests")
async def create_er(body: ERIn, db: Session = Depends(get_db), user: User = Depends(get_current_user)):
    tl = [{"at": datetime.now().isoformat(), "event": "Request created", "by": user.username}]
    r = EmergencyRequest(**body.model_dump(), status="requested", timeline=json.dumps(tl), created_by=user.id)
    db.add(r); db.flush()
    _audit(db, user.id, "create", "emergency_request", r.id, f"{body.emergency_type} @ {body.location}")
    _notify(db, "emergency", "🚨 New emergency request", f"#{r.id} {body.emergency_type} — {body.location}", role_target="emergency", related_id=r.id)
    _notify(db, "emergency", "🚨 New emergency request", f"#{r.id} {body.emergency_type} — {body.location}", role_target="admin", related_id=r.id)
    db.commit(); db.refresh(r)
    await hub.broadcast("emergency_request", {"request": _er(r)})
    return _er(r)


class ERPatch(BaseModel):
    status: Optional[str] = None
    ambulance_id: Optional[int] = None
    assigned_staff: Optional[str] = None
    notes: Optional[str] = None


@router.patch("/api/emergency/requests/{rid}")
async def patch_er(rid: int, body: ERPatch, db: Session = Depends(get_db), user: User = Depends(require_roles("admin", "emergency", "triage", "doctor"))):
    r = db.get(EmergencyRequest, rid)
    if not r:
        raise HTTPException(404, "Request not found")
    try:
        tl = json.loads(r.timeline or "[]")
    except Exception:
        tl = []
    if body.status:
        if body.status not in ("requested", "assigned", "en_route", "arrived", "transferred", "completed", "cancelled"):
            raise HTTPException(400, "Invalid status")
        r.status = body.status
        tl.append({"at": datetime.now().isoformat(), "event": f"Status → {body.status}", "by": user.username})
    if body.ambulance_id:
        amb = db.get(Ambulance, body.ambulance_id)
        if not amb:
            raise HTTPException(404, "Ambulance not found")
        r.ambulance_id = amb.id
        amb.status = "on_trip"
        tl.append({"at": datetime.now().isoformat(), "event": f"Ambulance {amb.vehicle_no} assigned", "by": user.username})
    if body.assigned_staff is not None:
        r.assigned_staff = body.assigned_staff
    if body.notes is not None:
        r.notes = body.notes
    r.timeline = json.dumps(tl)
    _audit(db, user.id, "update", "emergency_request", rid, body.model_dump_json()[:400])
    db.commit()
    await hub.broadcast("emergency_request", {"request": _er(r)})
    return _er(r)


@router.get("/api/emergency/ambulances")
def list_amb(db: Session = Depends(get_db), user: User = Depends(get_current_user)):
    return [{"id": a.id, "vehicle_no": a.vehicle_no, "status": a.status, "location": a.location, "crew": a.crew} for a in db.query(Ambulance).order_by(Ambulance.id).all()]


class AmbIn(BaseModel):
    vehicle_no: str
    status: str = "available"
    location: str = ""
    crew: str = ""


@router.post("/api/emergency/ambulances")
def create_amb(body: AmbIn, db: Session = Depends(get_db), user: User = Depends(require_roles("admin", "emergency"))):
    a = Ambulance(**body.model_dump())
    db.add(a); db.commit(); db.refresh(a)
    return {"id": a.id}


@router.patch("/api/emergency/ambulances/{aid}")
async def patch_amb(aid: int, body: dict, db: Session = Depends(get_db), user: User = Depends(require_roles("admin", "emergency"))):
    a = db.get(Ambulance, aid)
    if not a:
        raise HTTPException(404, "Ambulance not found")
    for k in ("status", "location", "crew"):
        if k in body:
            setattr(a, k, body[k])
    db.commit()
    await hub.broadcast("emergency_request", {"reason": "ambulance"})
    return {"ok": True}


# ================= ASSISTANT + ARTICLES =================
class ChatIn(BaseModel):
    message: str
    lang: str = "en"
    session_id: Optional[int] = None


@router.post("/api/assistant/chat")
def assistant_chat(body: ChatIn, db: Session = Depends(get_db), user: Optional[User] = Depends(get_current_user)):
    lang = body.lang if body.lang in ("en", "hi", "gu") else "en"
    out = chat_reply(body.message, lang)
    sid = body.session_id
    try:
        if user is not None:
            if not sid:
                s = ChatSession(user_id=user.id, lang=lang)
                db.add(s); db.flush()
                sid = s.id
            db.add(ChatMessage(session_id=sid, role="user", text=body.message[:2000]))
            db.add(ChatMessage(session_id=sid, role="assistant", text=out["reply"][:4000]))
            db.commit()
    except Exception:
        db.rollback()
    out["session_id"] = sid
    out["disclaimer"] = "General health education only — not a diagnosis. For emergencies call 108/112 (India)."
    return out


@router.get("/api/articles")
def list_articles(lang: str = "en", category: str = "", search: str = "", db: Session = Depends(get_db)):
    q = db.query(HealthArticle).filter_by(lang=lang if lang in ("en", "hi", "gu") else "en")
    if category:
        q = q.filter_by(category=category)
    if search:
        like = f"%{search}%"
        q = q.filter((HealthArticle.title.ilike(like)) | (HealthArticle.body.ilike(like)))
    return [{"slug": a.slug, "category": a.category, "title": a.title, "body": a.body, "source": a.source} for a in q.order_by(HealthArticle.id).limit(100)]


@router.get("/api/articles/{slug}")
def get_article(slug: str, lang: str = "en", db_session: Session = Depends(get_db)):
    a = db_session.query(HealthArticle).filter_by(slug=slug, lang=lang).first() or db_session.query(HealthArticle).filter_by(slug=slug).first()
    if not a:
        raise HTTPException(404, "Article not found")
    return {"slug": a.slug, "category": a.category, "title": a.title, "body": a.body, "source": a.source, "lang": a.lang}


@router.get("/api/nutrition")
def list_nutrition(lang: str = "en", group: str = "", db: Session = Depends(get_db)):
    q = db.query(NutritionResource).filter_by(lang=lang if lang in ("en", "hi", "gu") else "en")
    if group:
        q = q.filter_by(group=group)
    return [{"id": n.id, "title": n.title, "group": n.group, "body": n.body, "source": n.source} for n in q.order_by(NutritionResource.id).limit(60)]


@router.get("/api/nutrition/meals")
def meal_ideas():
    return [{"en": m[0], "en_detail": m[1], "hi": m[2], "hi_detail": m[3]} for m in MEALS]


# ================= WELLNESS =================
class MoodIn(BaseModel):
    mood: int
    note: str = ""
    sleep_hours: Optional[float] = None


@router.get("/api/wellness/logs")
def my_logs(db: Session = Depends(get_db), user: User = Depends(get_current_user)):
    rows = db.query(WellnessLog).filter_by(user_id=user.id).order_by(WellnessLog.id.desc()).limit(120).all()
    return [{"id": l.id, "mood": l.mood, "note": l.note, "sleep_hours": l.sleep_hours, "at": l.created_at.isoformat() if l.created_at else None} for l in rows]


@router.post("/api/wellness/logs")
def add_log(body: MoodIn, db: Session = Depends(get_db), user: User = Depends(get_current_user)):
    if body.mood < 1 or body.mood > 5:
        raise HTTPException(400, "Mood must be 1–5")
    l = WellnessLog(user_id=user.id, mood=body.mood, note=body.note[:1000], sleep_hours=body.sleep_hours)
    db.add(l); db.commit(); db.refresh(l)
    return {"id": l.id}


@router.delete("/api/wellness/logs/{lid}")
def del_log(lid: int, db: Session = Depends(get_db), user: User = Depends(get_current_user)):
    l = db.get(WellnessLog, lid)
    if not l or (l.user_id != user.id and user.role != "admin"):
        raise HTTPException(404, "Log not found")
    db.delete(l); db.commit()
    return {"ok": True}


@router.get("/api/wellness/summary")
def wellness_summary(db: Session = Depends(get_db), user: User = Depends(get_current_user)):
    since = datetime.now() - timedelta(days=14)
    rows = db.query(WellnessLog).filter(WellnessLog.user_id == user.id, WellnessLog.created_at >= since).order_by(WellnessLog.created_at).all()
    moods = [r.mood for r in rows]
    return {"count_14d": len(rows), "avg_mood": round(sum(moods)/len(moods), 1) if moods else None,
            "trend": [{"at": r.created_at.isoformat()[:10], "mood": r.mood} for r in rows]}


# ================= MEDICINES / SERVICES =================
@router.get("/api/medicines")
def list_meds(search: str = "", category: str = "", db: Session = Depends(get_db)):
    q = db.query(Medicine)
    if search:
        like = f"%{search}%"
        q = q.filter((Medicine.name.ilike(like)) | (Medicine.generic.ilike(like)) | (Medicine.brands.ilike(like)) | (Medicine.uses.ilike(like)))
    if category:
        q = q.filter_by(category=category)
    rows = q.order_by(Medicine.name).limit(200).all()
    return [{"id": m.id, "name": m.name, "generic": m.generic, "brands": m.brands, "category": m.category, "uses": m.uses, "safety": m.safety, "warnings": m.warnings, "manufacturer": m.manufacturer, "source": m.source} for m in rows]


@router.get("/api/medicines/{mid}")
def get_med(mid: int, db: Session = Depends(get_db)):
    m = db.get(Medicine, mid)
    if not m:
        raise HTTPException(404, "Medicine not found")
    return {"id": m.id, "name": m.name, "generic": m.generic, "brands": m.brands, "category": m.category, "uses": m.uses, "safety": m.safety, "warnings": m.warnings, "manufacturer": m.manufacturer, "source": m.source}


@router.get("/api/services")
def list_services(type: str = "", search: str = "", db: Session = Depends(get_db)):
    q = db.query(ServiceDir)
    if type:
        q = q.filter_by(type=type)
    if search:
        like = f"%{search}%"
        q = q.filter((ServiceDir.name.ilike(like)) | (ServiceDir.place.ilike(like)) | (ServiceDir.services.ilike(like)))
    return [{"id": s.id, "name": s.name, "type": s.type, "place": s.place, "phone": s.phone, "hours": s.hours, "services": s.services} for s in q.order_by(ServiceDir.id).limit(200)]


# ================= BILLING / BEDS =================
class InvIn(BaseModel):
    patient_id: int
    visit_id: Optional[int] = None
    items: List[dict] = []


@router.get("/api/billing/invoices")
def list_inv(patient_id: Optional[int] = None, status: str = "", db: Session = Depends(get_db), user: User = Depends(get_current_user)):
    q = db.query(Invoice)
    if user.role == "patient":
        myp = [p.id for p in db.query(Patient).filter_by(user_id=user.id).all()]
        q = q.filter(Invoice.patient_id.in_(myp or [-1]))
    if patient_id:
        q = q.filter_by(patient_id=patient_id)
    if status:
        q = q.filter_by(status=status)
    rows = q.order_by(Invoice.id.desc()).limit(300).all()
    out = []
    for i in rows:
        try:
            items = json.loads(i.items or "[]")
        except Exception:
            items = []
        out.append({"id": i.id, "patient_id": i.patient_id, "patient_name": i.patient.full_name if i.patient else "", "visit_id": i.visit_id, "items": items, "total": i.total, "status": i.status, "at": i.created_at.isoformat() if i.created_at else None})
    return out


@router.post("/api/billing/invoices")
def create_inv(body: InvIn, db: Session = Depends(get_db), user: User = Depends(require_roles("admin", "triage"))):
    total = sum(float(it.get("amount", 0)) for it in body.items)
    i = Invoice(patient_id=body.patient_id, visit_id=body.visit_id, items=json.dumps(body.items), total=total, status="unpaid")
    db.add(i)
    _audit(db, user.id, "create", "invoice", None, f"patient {body.patient_id} total {total}")
    db.commit(); db.refresh(i)
    return {"id": i.id, "total": total}


@router.patch("/api/billing/invoices/{iid}")
def pay_inv(iid: int, body: dict, db: Session = Depends(get_db), user: User = Depends(require_roles("admin", "triage"))):
    i = db.get(Invoice, iid)
    if not i:
        raise HTTPException(404, "Invoice not found")
    if body.get("status") in ("unpaid", "paid", "waived"):
        i.status = body["status"]
    _audit(db, user.id, "update", "invoice", iid, i.status)
    db.commit()
    return {"ok": True}


@router.get("/api/beds")
def bed_summary(db: Session = Depends(get_db), user: User = Depends(get_current_user)):
    wards = db.query(Bed.ward, Bed.status, func.count(Bed.id)).group_by(Bed.ward, Bed.status).all()
    by_ward = {}
    for w, s, c in wards:
        by_ward.setdefault(w, {"ward": w, "available": 0, "occupied": 0, "reserved": 0, "maintenance": 0, "total": 0})
        by_ward[w][s] = c
        by_ward[w]["total"] += c
    return list(by_ward.values())


@router.patch("/api/beds/{bid}")
async def patch_bed(bid: int, body: dict, db: Session = Depends(get_db), user: User = Depends(require_roles("admin", "triage", "doctor"))):
    b = db.get(Bed, bid)
    if not b:
        raise HTTPException(404, "Bed not found")
    if body.get("status") in ("available", "occupied", "reserved", "maintenance"):
        b.status = body["status"]
    db.commit()
    await hub.broadcast("resources_updated", {"id": bid})
    return {"ok": True}


# ================= GLOBAL SEARCH =================
@router.get("/api/search")
def global_search(q: str = "", db: Session = Depends(get_db), user: User = Depends(get_current_user)):
    q = (q or "").strip()
    if len(q) < 2:
        return {"patients": [], "doctors": [], "departments": [], "appointments": [], "articles": [], "medicines": [], "services": []}
    like = f"%{q}%"
    pats = db.query(Patient).filter((Patient.full_name.ilike(like)) | (Patient.patient_code.ilike(like)) | (Patient.phone.ilike(like))).limit(8).all()
    if user.role == "patient":
        pats = [p for p in pats if p.user_id == user.id]
    docs = db.query(Doctor).filter(Doctor.full_name.ilike(like)).limit(8).all()
    deps = db.query(Department).filter(Department.name.ilike(like)).limit(8).all()
    appts = db.query(Appointment).options(joinedload(Appointment.patient)).join(Patient).filter(Patient.full_name.ilike(like)).limit(8).all() if user.role != "patient" else []
    arts = db.query(HealthArticle).filter(HealthArticle.title.ilike(like)).limit(8).all()
    meds = db.query(Medicine).filter((Medicine.name.ilike(like)) | (Medicine.generic.ilike(like))).limit(8).all()
    svcs = db.query(ServiceDir).filter((ServiceDir.name.ilike(like)) | (ServiceDir.place.ilike(like))).limit(8).all()
    return {
        "patients": [{"id": p.id, "label": f"{p.patient_code} — {p.full_name}"} for p in pats],
        "doctors": [{"id": d.id, "label": d.full_name} for d in docs],
        "departments": [{"id": d.id, "label": d.name} for d in deps],
        "appointments": [{"id": a.id, "label": f"#{a.id} {a.patient.full_name if a.patient else ''} — {a.date}"} for a in appts],
        "articles": [{"id": a.slug, "label": a.title} for a in arts],
        "medicines": [{"id": m.id, "label": f"{m.name} ({m.generic})"} for m in meds],
        "services": [{"id": s.id, "label": f"{s.name} — {s.place}"} for s in svcs],
    }


# ================= SETTINGS =================
@router.get("/api/settings")
def get_settings(db: Session = Depends(get_db), user: User = Depends(get_current_user)):
    s = db.get(UserSetting, user.id)
    return {"theme": s.theme if s else "light", "language": s.language if s else "en"}


@router.patch("/api/settings")
def patch_settings(body: dict, db: Session = Depends(get_db), user: User = Depends(get_current_user)):
    s = db.get(UserSetting, user.id)
    if not s:
        s = UserSetting(user_id=user.id)
        db.add(s)
    if body.get("theme") in ("light", "dark"):
        s.theme = body["theme"]
    if body.get("language") in ("en", "hi", "gu"):
        s.language = body["language"]
    db.commit()
    return {"theme": s.theme, "language": s.language}


# ================= ECOSYSTEM ANALYTICS =================
@router.get("/api/analytics/ecosystem")
def eco_stats(db: Session = Depends(get_db), user: User = Depends(get_current_user)):
    from .models import HospitalResource
    today = date.today()
    appts_today = db.query(Appointment).filter_by(date=today).count()
    appts_up = db.query(Appointment).filter(Appointment.date >= today, Appointment.status.in_(["requested", "confirmed", "rescheduled"])).count()
    blood_units = db.query(func.coalesce(func.sum(BloodInventory.units_available), 0)).scalar() or 0
    blood_open = db.query(BloodRequest).filter_by(status="requested").count()
    er_active = db.query(EmergencyRequest).filter(EmergencyRequest.status.in_(["requested", "assigned", "en_route", "arrived", "transferred"])).count()
    amb_avail = db.query(Ambulance).filter_by(status="available").count()
    beds = db.query(Bed).all()
    bed_tot, bed_occ = len(beds), sum(1 for b in beds if b.status == "occupied")
    inv_out = db.query(func.coalesce(func.sum(Invoice.total), 0)).filter_by(status="unpaid").scalar() or 0
    donors = db.query(BloodDonor).count()
    return {"appointments_today": appts_today, "appointments_upcoming": appts_up, "blood_units": int(blood_units),
            "blood_open_requests": blood_open, "emergencies_active": er_active, "ambulances_available": amb_avail,
            "beds_total": bed_tot, "beds_occupied": bed_occ, "bed_utilization": round(100 * bed_occ / bed_tot, 1) if bed_tot else 0,
            "unpaid_billing": round(float(inv_out), 2), "donors": donors}


@router.get("/api/analytics/ml/demand")
def demand_forecast(db: Session = Depends(get_db), user: User = Depends(get_current_user)):
    """7-day arrival forecast = 7-day moving average (transparent operational heuristic)."""
    import pandas as pd
    daily = []
    for i in range(60, 0, -1):
        d = date.today() - timedelta(days=i)
        s = datetime(d.year, d.month, d.day)
        e = s + timedelta(days=1)
        daily.append(db.query(Visit).filter(Visit.arrival_time >= s, Visit.arrival_time < e).count())
    s = pd.Series(daily, dtype=float)
    ma = float(s.tail(7).mean())
    trend = float(s.tail(7).mean() - s.tail(14).head(7).mean())
    forecast = [max(0, round(ma + trend * (i / 7))) for i in range(1, 8)]
    return {"history_last_14": daily[-14:], "forecast_next_7": [{"day": f"Day+{i+1}", "estimate": v} for i, v in enumerate(forecast)],
            "method": "7-day moving average + trend", "note": "Operational estimate only — not a clinical prediction."}


@router.get("/api/analytics/utilization")
def utilization(db: Session = Depends(get_db), user: User = Depends(get_current_user)):
    from .models import HospitalResource
    docs_tot = db.query(Doctor).count()
    docs_av = db.query(Doctor).filter_by(available=True).count()
    wl = db.query(Department.name, func.count(Visit.id)).join(Visit, Visit.department_id == Department.id).filter(Visit.status.in_(["waiting", "called", "under_assessment", "registered"])).group_by(Department.name).all()
    res = db.query(HospitalResource).all()
    amb = db.query(Ambulance.status, func.count(Ambulance.id)).group_by(Ambulance.status).all()
    return {"doctors": {"total": docs_tot, "available": docs_av},
            "dept_queue": [{"name": n, "value": c} for n, c in wl],
            "resources": [{"ward": r.ward_name, "total": r.total, "occupied": r.occupied} for r in res],
            "ambulances": [{"name": s, "value": c} for s, c in amb]}


@router.get("/api/analytics/data-quality")
def data_quality(db: Session = Depends(get_db), user: User = Depends(require_roles("admin", "doctor"))):
    """Pandas-based data-quality snapshot over visits + triage (B.Sc. demo)."""
    import pandas as pd
    from .models import TriageAssessment
    visits = db.query(Visit).order_by(Visit.id.desc()).limit(5000).all()
    df = pd.DataFrame([{"priority": v.priority, "status": v.status, "dept": v.department_id, "waiting": v.waiting_minutes} for v in visits])
    tri = db.query(TriageAssessment).order_by(TriageAssessment.id.desc()).limit(5000).all()
    tv = pd.DataFrame([{"temp": t.temperature, "hr": t.heart_rate, "sbp": t.sbp, "spo2": t.spo2, "rr": t.respiratory_rate} for t in tri])
    return {
        "visits_rows": len(df),
        "missing_pct": {c: round(float(df[c].isna().mean() * 100), 1) for c in df.columns} if len(df) else {},
        "vitals_missing_pct": {c: round(float(tv[c].isna().mean() * 100), 1) for c in tv.columns} if len(tv) else {},
        "vitals_describe": tv.describe().round(1).to_dict() if len(tv) else {},
        "duplicates": int(df.duplicated().sum()) if len(df) else 0,
        "steps": ["Collection (API/DB)", "Validation (Pydantic)", "Missing-value analysis", "Cleaning (median impute demo)", "EDA", "Features", "ML", "Evaluation", "Dashboard"],
    }


# ================= RISK ML LAB (educational) =================
_risk_models = {}
_risk_metrics = {}


def _risk_dataset(n: int = 2000):
    import numpy as np
    rng = np.random.default_rng(42)
    age = rng.integers(20, 80, n)
    bmi = rng.normal(26, 5, n).clip(15, 45)
    glucose = rng.normal(110, 30, n).clip(70, 260)
    bp = rng.normal(125, 18, n).clip(90, 200)
    fam = rng.integers(0, 2, n)
    smoker = rng.integers(0, 2, n)
    logit = -8.0 + 0.04 * age + 0.08 * bmi + 0.02 * glucose + 0.01 * bp + 0.9 * fam + 0.6 * smoker
    prob = 1 / (1 + np.exp(-logit))
    y = (rng.random(n) < prob).astype(int)
    import pandas as pd
    X = pd.DataFrame({"age": age, "bmi": bmi.round(1), "glucose": glucose.round(0), "bp": bp.round(0), "family_history": fam, "smoker": smoker})
    return X, y


@router.post("/api/ml/risk/train")
def risk_train(body: dict = {}, db: Session = Depends(get_db), user: User = Depends(require_roles("admin", "doctor"))):
    from sklearn.linear_model import LogisticRegression
    from sklearn.ensemble import RandomForestClassifier
    from sklearn.tree import DecisionTreeClassifier
    from sklearn.model_selection import train_test_split
    from sklearn.metrics import accuracy_score, precision_score, recall_score, f1_score, confusion_matrix, roc_auc_score
    which = (body.get("model") or "logistic")
    X, y = _risk_dataset()
    Xtr, Xte, ytr, yte = train_test_split(X, y, test_size=0.25, random_state=42, stratify=y)
    m = {"logistic": LogisticRegression(max_iter=500), "rf": RandomForestClassifier(n_estimators=150, random_state=42, n_jobs=-1), "tree": DecisionTreeClassifier(max_depth=6, random_state=42)}.get(which, LogisticRegression(max_iter=500))
    m.fit(Xtr, ytr)
    pred = m.predict(Xte)
    proba = m.predict_proba(Xte)[:, 1]
    met = {"model": which, "n": len(X), "accuracy": round(float(accuracy_score(yte, pred)), 3),
           "precision": round(float(precision_score(yte, pred, zero_division=0)), 3),
           "recall": round(float(recall_score(yte, pred, zero_division=0)), 3),
           "f1": round(float(f1_score(yte, pred, zero_division=0)), 3),
           "roc_auc": round(float(roc_auc_score(yte, proba)), 3),
           "confusion_matrix": confusion_matrix(yte, pred).tolist(),
           "features": list(X.columns),
           "note": "Model-based risk estimate for EDUCATIONAL/RESEARCH purposes only — not a diagnosis."}
    _risk_models[which] = m
    _risk_metrics[which] = met
    _audit(db, user.id, "ml_train", "risk_model", None, which)
    db.commit()
    return met


@router.get("/api/ml/risk/metrics")
def risk_metrics(user: User = Depends(get_current_user)):
    return _risk_metrics or {"status": "not-trained", "hint": "POST /api/ml/risk/train with {model: logistic|rf|tree}"}


class RiskIn(BaseModel):
    age: float = 45
    bmi: float = 27
    glucose: float = 120
    bp: float = 130
    family_history: int = 0
    smoker: int = 0
    model: str = "logistic"


@router.post("/api/ml/risk/predict")
def risk_predict(body: RiskIn, db: Session = Depends(get_db), user: User = Depends(get_current_user)):
    import pandas as pd
    m = _risk_models.get(body.model)
    if m is None:
        raise HTTPException(400, "Model not trained yet — train it first (admin/doctor)")
    X = pd.DataFrame([{k: getattr(body, k) for k in ("age", "bmi", "glucose", "bp", "family_history", "smoker")}])
    p = float(m.predict_proba(X)[0, 1])
    band = "Lower estimated risk" if p < 0.33 else ("Moderate estimated risk" if p < 0.66 else "Higher estimated risk")
    db.add(MlPrediction(model=f"risk_{body.model}", input_json=X.to_json(), output=f"{band} ({p:.2f})", created_by=user.id))
    db.commit()
    return {"probability": round(p, 3), "band": band, "model": body.model,
            "disclaimer": "Educational model estimate on synthetic data — NOT a medical diagnosis. Consult a qualified professional."}


# ================= EXTRA REPORTS =================
def _appt_rows(db, days=7):
    since = datetime.now() - timedelta(days=days)
    rows = db.query(Appointment).options(joinedload(Appointment.patient), joinedload(Appointment.doctor)).filter(Appointment.created_at >= since).order_by(Appointment.date).limit(1000).all()
    return [{"id": a.id, "patient": a.patient.full_name if a.patient else "", "doctor": a.doctor.full_name if a.doctor else "",
             "date": a.date.isoformat() if a.date else "", "slot": a.time_slot, "status": a.status, "reason": (a.reason or "")[:40]} for a in rows]


@router.get("/api/reports/appointments")
def rep_appts(format: str = "json", days: int = 7, db: Session = Depends(get_db), user: User = Depends(require_roles("admin", "doctor"))):
    rows = _appt_rows(db, days)
    return _export(rows, ["id", "patient", "doctor", "date", "slot", "status", "reason"], "appointment_report", format)


@router.get("/api/reports/blood")
def rep_blood(format: str = "json", db: Session = Depends(get_db), user: User = Depends(require_roles("admin", "blood", "doctor"))):
    inv = [{"group": r.blood_group, "available": r.units_available, "reserved": r.units_reserved} for r in db.query(BloodInventory).all()]
    reqs = db.query(BloodRequest).order_by(BloodRequest.id.desc()).limit(500).all()
    rows = [{"id": r.id, "unit": r.hospital_unit, "group": r.blood_group, "units": r.units, "priority": r.priority, "status": r.status} for r in reqs]
    if format == "json":
        return {"inventory": inv, "requests": rows, "synthetic_data": True}
    return _export(rows, ["id", "unit", "group", "units", "priority", "status"], "blood_bank_report", format)


@router.get("/api/reports/emergency")
def rep_er(format: str = "json", db: Session = Depends(get_db), user: User = Depends(require_roles("admin", "emergency", "doctor"))):
    rows = db.query(EmergencyRequest).order_by(EmergencyRequest.id.desc()).limit(500).all()
    data = [{"id": r.id, "patient": r.patient_name, "type": r.emergency_type, "location": (r.location or "")[:30], "priority": r.priority, "status": r.status, "ambulance": r.ambulance.vehicle_no if r.ambulance else ""} for r in rows]
    return _export(data, ["id", "patient", "type", "location", "priority", "status", "ambulance"], "emergency_report", format)
